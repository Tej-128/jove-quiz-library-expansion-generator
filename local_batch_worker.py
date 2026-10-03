from __future__ import annotations

import argparse
import csv
import json
import os
import shutil
import threading
import time
import traceback
import uuid
from concurrent.futures import FIRST_COMPLETED, ThreadPoolExecutor, wait
from pathlib import Path
from typing import Any

import openpyxl

from input_parser import bundle_lessons, records_from_directory
from local_batch_runtime import CONFIG_FILENAME, STATUS_FILENAME, STOP_FILENAME, read_json
from pipeline import PIPELINE_VERSION, process_lesson
from quiz_generator import GENERATED_TYPES, PER_TYPE, TOTAL_GENERATED

RATE_LIMIT_MARKERS = (
    "429",
    "rate limit",
    "rate_limit",
    "too many requests",
    "requests per min",
    "tokens per min",
)


def atomic_json(path: Path, payload: dict[str, Any]) -> None:
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    os.replace(tmp, path)


def safe_csv_value(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\r", " ").replace("\n", " ").strip()


def write_progress_csv(path: Path, lesson_ids: list[str], rows: dict[str, dict[str, Any]]) -> None:
    fieldnames = [
        "Lesson ID",
        "Status",
        "Attempt",
        "Lesson Title",
        "Chapter",
        "Existing Questions",
        "New Questions",
        "Total Questions",
        "New Review Flags",
        "New Red Flags",
        "QA Replacements",
        "QA Repair Rounds",
        "Elapsed Seconds",
        "Output File",
        "Last Error",
    ]
    tmp = path.with_suffix(path.suffix + ".tmp")
    with tmp.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for lesson_id in lesson_ids:
            raw = rows.get(lesson_id, {})
            writer.writerow(
                {
                    "Lesson ID": lesson_id,
                    "Status": safe_csv_value(raw.get("status", "Pending")),
                    "Attempt": safe_csv_value(raw.get("attempt", 0)),
                    "Lesson Title": safe_csv_value(raw.get("lesson_title", "")),
                    "Chapter": safe_csv_value(raw.get("chapter", "")),
                    "Existing Questions": safe_csv_value(raw.get("existing_questions", "")),
                    "New Questions": safe_csv_value(raw.get("new_questions", "")),
                    "Total Questions": safe_csv_value(raw.get("total_questions", "")),
                    "New Review Flags": safe_csv_value(raw.get("review_flags", "")),
                    "New Red Flags": safe_csv_value(raw.get("red_flags", "")),
                    "QA Replacements": safe_csv_value(raw.get("qa_replacements", "")),
                    "QA Repair Rounds": safe_csv_value(raw.get("qa_repair_rounds", "")),
                    "Elapsed Seconds": safe_csv_value(raw.get("elapsed_seconds", "")),
                    "Output File": safe_csv_value(raw.get("output_file", "")),
                    "Last Error": safe_csv_value(raw.get("last_error", "")),
                }
            )
    os.replace(tmp, path)


def load_existing_rows(path: Path) -> dict[str, dict[str, Any]]:
    if not path.exists():
        return {}
    rows: dict[str, dict[str, Any]] = {}
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            for row in csv.DictReader(handle):
                lesson_id = str(row.get("Lesson ID", "")).strip()
                if not lesson_id:
                    continue
                rows[lesson_id] = {
                    "status": row.get("Status", "Pending"),
                    "attempt": int(row.get("Attempt", "0") or 0),
                    "lesson_title": row.get("Lesson Title", ""),
                    "chapter": row.get("Chapter", ""),
                    "existing_questions": row.get("Existing Questions", ""),
                    "new_questions": row.get("New Questions", ""),
                    "total_questions": row.get("Total Questions", ""),
                    "review_flags": row.get("New Review Flags", ""),
                    "red_flags": row.get("New Red Flags", ""),
                    "qa_replacements": row.get("QA Replacements", ""),
                    "qa_repair_rounds": row.get("QA Repair Rounds", ""),
                    "elapsed_seconds": row.get("Elapsed Seconds", ""),
                    "output_file": row.get("Output File", ""),
                    "last_error": row.get("Last Error", ""),
                }
    except Exception:
        return {}
    return rows


def is_rate_limit_error(message: str) -> bool:
    lower = str(message or "").lower()
    return any(marker in lower for marker in RATE_LIMIT_MARKERS)


def validate_output(path: str, report: dict[str, Any]) -> None:
    if int(report.get("generated_questions", -1)) != TOTAL_GENERATED:
        raise RuntimeError(
            f"Output validation failed: expected {TOTAL_GENERATED} generated questions, "
            f"got {report.get('generated_questions')}."
        )

    type_counts = report.get("generated_type_counts", {}) or {}
    wrong = {
        qtype: type_counts.get(qtype, 0)
        for qtype in GENERATED_TYPES
        if int(type_counts.get(qtype, 0) or 0) != PER_TYPE
    }
    if wrong:
        raise RuntimeError(f"Output validation failed: generated type counts are incorrect: {wrong}")

    output = Path(path)
    if not output.is_file() or output.stat().st_size < 1000:
        raise RuntimeError("Output validation failed: workbook file is missing or unexpectedly small.")

    wb = openpyxl.load_workbook(output, read_only=True, data_only=False)
    try:
        if "Quiz" not in wb.sheetnames or "QA Summary" not in wb.sheetnames:
            raise RuntimeError("Output validation failed: required workbook sheets are missing.")
    finally:
        wb.close()


def lesson_task(bundle, config: dict[str, Any], job_dir: Path, attempt: int) -> dict[str, Any]:
    lesson_id = bundle.lesson_id
    started = time.monotonic()
    work_dir = job_dir / "work" / f"{lesson_id}_attempt{attempt}_{uuid.uuid4().hex[:6]}"
    work_dir.mkdir(parents=True, exist_ok=True)

    try:
        output_path, report = process_lesson(
            bundle,
            subject=config["subject"],
            api_key=os.environ["OPENAI_API_KEY"],
            model=config["model"],
            enable_ai_accuracy_qa=bool(config.get("enable_ai_accuracy_qa", True)),
            output_dir=str(work_dir),
            progress_callback=None,
        )
        validate_output(output_path, report)

        output_root = Path(config["output_dir"])
        output_root.mkdir(parents=True, exist_ok=True)
        final_path = output_root / Path(output_path).name

        # Only complete, validated workbooks are exposed to the user output folder.
        os.replace(output_path, final_path)

        return {
            "lesson_id": lesson_id,
            "success": True,
            "attempt": attempt,
            "elapsed_seconds": int(time.monotonic() - started),
            "output_file": str(final_path),
            "report": report,
        }
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)


def main(job_dir: str) -> int:
    job_path = Path(job_dir).resolve()
    config = read_json(job_path / CONFIG_FILENAME, {})
    if not config:
        raise RuntimeError(f"Job configuration is missing: {job_path / CONFIG_FILENAME}")

    api_key = os.environ.get("OPENAI_API_KEY", "").strip()
    if not api_key:
        raise RuntimeError("OPENAI_API_KEY is missing from the local worker environment.")

    if config.get("build_version") != PIPELINE_VERSION:
        raise RuntimeError(
            f"Build mismatch: job={config.get('build_version')} current={PIPELINE_VERSION}. "
            "Create a new job with the current build."
        )

    lesson_ids = [str(value) for value in config.get("lesson_ids", [])]
    input_dir = Path(config["input_dir"])
    output_dir = Path(config["output_dir"])
    progress_csv = output_dir / f"generation_progress_{config['job_id']}.csv"
    status_path = job_path / STATUS_FILENAME
    rows = load_existing_rows(progress_csv)

    records = records_from_directory(str(input_dir))
    bundles = bundle_lessons(records)
    bundle_by_id = {
        bundle.lesson_id: bundle
        for bundle in bundles
        if bundle.lesson_id != "UNASSIGNED"
    }

    for lesson_id in lesson_ids:
        rows.setdefault(
            lesson_id,
            {
                "status": "Pending",
                "attempt": 0,
                "last_error": "",
            },
        )

    completed_ids = {
        lesson_id
        for lesson_id, row in rows.items()
        if row.get("status") == "Completed"
        and row.get("output_file")
        and Path(str(row.get("output_file"))).is_file()
    }

    missing_bundle_ids = [
        lesson_id
        for lesson_id in lesson_ids
        if lesson_id not in bundle_by_id
    ]
    for lesson_id in missing_bundle_ids:
        rows[lesson_id].update(
            {
                "status": "Failed",
                "last_error": "Lesson source bundle is missing or no longer resolves to PageText + Transcript/CC + Quiz.",
            }
        )

    max_attempts = max(1, int(config.get("max_lesson_attempts", 3) or 3))
    requested_workers = max(1, min(int(config.get("workers", 12) or 12), 16))
    current_workers = requested_workers
    lock = threading.Lock()

    recent_completed: list[str] = []
    recent_failed: list[str] = []
    start_time = int(time.time())

    def persist(state: str, message: str, in_progress: int = 0) -> None:
        with lock:
            completed = sum(
                1
                for lesson_id in lesson_ids
                if rows.get(lesson_id, {}).get("status") == "Completed"
            )
            failed = sum(
                1
                for lesson_id in lesson_ids
                if rows.get(lesson_id, {}).get("status") == "Failed"
            )
            payload = {
                "job_id": config["job_id"],
                "state": state,
                "total": len(lesson_ids),
                "completed": completed,
                "failed": failed,
                "in_progress": in_progress,
                "workers_active": current_workers,
                "workers_requested": requested_workers,
                "recent_completed": recent_completed[-10:],
                "recent_failed": recent_failed[-10:],
                "started_at": start_time,
                "updated_at": int(time.time()),
                "message": message,
                "progress_csv": str(progress_csv),
                "output_dir": str(output_dir),
            }
            atomic_json(status_path, payload)
            write_progress_csv(progress_csv, lesson_ids, rows)

    persist(
        "running",
        f"Local worker started with {current_workers} parallel lesson workers.",
        0,
    )

    pending = [
        lesson_id
        for lesson_id in lesson_ids
        if lesson_id not in completed_ids
        and lesson_id in bundle_by_id
    ]

    for attempt_round in range(1, max_attempts + 1):
        if not pending:
            break

        if (job_path / STOP_FILENAME).exists():
            persist(
                "paused",
                "Stop requested. Completed files are preserved; resume will continue unfinished lessons.",
                0,
            )
            return 0

        round_ids = [
            lesson_id
            for lesson_id in pending
            if int(rows[lesson_id].get("attempt", 0) or 0) < max_attempts
        ]
        pending = []
        if not round_ids:
            break

        rate_limited_failures = 0
        failures_this_round: list[str] = []

        with ThreadPoolExecutor(
            max_workers=min(current_workers, len(round_ids)),
            thread_name_prefix="jove-lesson",
        ) as executor:
            future_map = {}
            for lesson_id in round_ids:
                if (job_path / STOP_FILENAME).exists():
                    break

                attempt = int(rows[lesson_id].get("attempt", 0) or 0) + 1
                rows[lesson_id].update(
                    {
                        "status": "Running",
                        "attempt": attempt,
                        "last_error": "",
                    }
                )
                future = executor.submit(
                    lesson_task,
                    bundle_by_id[lesson_id],
                    config,
                    job_path,
                    attempt,
                )
                future_map[future] = (lesson_id, attempt)

            persist(
                "running",
                f"Processing {len(future_map)} lesson(s) in retry round {attempt_round}/{max_attempts}.",
                len(future_map),
            )

            unfinished = set(future_map)
            while unfinished:
                done, unfinished = wait(
                    unfinished,
                    timeout=1.0,
                    return_when=FIRST_COMPLETED,
                )
                for future in done:
                    lesson_id, attempt = future_map[future]
                    try:
                        result = future.result()
                        report = result["report"]
                        rows[lesson_id].update(
                            {
                                "status": "Completed",
                                "attempt": attempt,
                                "lesson_title": report.get("lesson_title", ""),
                                "chapter": report.get("chapter_name", ""),
                                "existing_questions": report.get("existing_questions", ""),
                                "new_questions": report.get("generated_questions", ""),
                                "total_questions": report.get("total_questions", ""),
                                "review_flags": report.get("generated_review_count", ""),
                                "red_flags": report.get("generated_fail_count", ""),
                                "qa_replacements": report.get("qa_replacements", ""),
                                "qa_repair_rounds": report.get("qa_repair_rounds", ""),
                                "elapsed_seconds": result.get("elapsed_seconds", ""),
                                "output_file": result.get("output_file", ""),
                                "last_error": "",
                            }
                        )
                        recent_completed.append(lesson_id)
                    except Exception as exc:
                        error_text = f"{type(exc).__name__}: {exc}"
                        rows[lesson_id].update(
                            {
                                "status": "Retrying" if attempt < max_attempts else "Failed",
                                "attempt": attempt,
                                "last_error": error_text,
                            }
                        )
                        recent_failed.append(lesson_id)
                        failures_this_round.append(lesson_id)
                        if is_rate_limit_error(error_text):
                            rate_limited_failures += 1

                    completed_now = sum(
                        1
                        for x in lesson_ids
                        if rows[x].get("status") == "Completed"
                    )
                    persist(
                        "running",
                        f"Completed {completed_now}/{len(lesson_ids)} lessons.",
                        len(unfinished),
                    )

        pending = [
            lesson_id
            for lesson_id in failures_this_round
            if int(rows[lesson_id].get("attempt", 0) or 0) < max_attempts
        ]

        if pending and rate_limited_failures:
            previous_workers = current_workers
            current_workers = max(4, current_workers // 2)
            backoff = min(120, 30 * attempt_round)
            persist(
                "running",
                f"Rate limiting detected. Reducing workers {previous_workers}->{current_workers} and backing off {backoff}s before retries.",
                0,
            )
            time.sleep(backoff)
        elif pending:
            backoff = min(45, 10 * attempt_round)
            persist(
                "running",
                f"Retrying {len(pending)} lesson(s) after {backoff}s.",
                0,
            )
            time.sleep(backoff)

    for lesson_id in lesson_ids:
        if rows[lesson_id].get("status") in {"Pending", "Running", "Retrying"}:
            rows[lesson_id]["status"] = "Failed"
            if not rows[lesson_id].get("last_error"):
                rows[lesson_id]["last_error"] = "Lesson did not complete within the retry budget."

    failed_count = sum(
        1
        for lesson_id in lesson_ids
        if rows[lesson_id].get("status") == "Failed"
    )
    final_state = "completed" if failed_count == 0 else "completed_with_failures"

    persist(
        final_state,
        (
            f"Batch finished successfully: {len(lesson_ids)} lesson(s) completed."
            if failed_count == 0
            else f"Batch finished with {failed_count} failed lesson(s). Completed outputs remain available; failed lessons can be resumed or retried."
        ),
        0,
    )
    return 0


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--job-dir", required=True)
    args = parser.parse_args()

    try:
        raise SystemExit(main(args.job_dir))
    except SystemExit:
        raise
    except Exception:
        job_path = Path(args.job_dir).resolve()
        fatal_path = job_path / "fatal_error.log"
        fatal_path.write_text(traceback.format_exc(), encoding="utf-8")

        status_path = job_path / STATUS_FILENAME
        current = read_json(status_path, {})
        current.update(
            {
                "state": "worker_crashed",
                "updated_at": int(time.time()),
                "message": "Worker process crashed; the supervisor will restart it automatically.",
            }
        )
        atomic_json(status_path, current)
        raise
